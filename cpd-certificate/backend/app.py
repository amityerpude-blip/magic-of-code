import os,re,smtplib,subprocess,uuid
from pathlib import Path
from email.message import EmailMessage
from flask import Flask,request,jsonify,send_from_directory
from flask_cors import CORS
from pptx import Presentation
ROOT=Path(__file__).parent; TEMPLATE=ROOT/'cpd_template.pptx'; OUT=ROOT/'generated'; OUT.mkdir(exist_ok=True)
app=Flask(__name__); CORS(app,resources={r'/api/*':{'origins':os.getenv('FRONTEND_ORIGIN','*')}})
FIELDS=['email','name','hname','post','hpost','training','htraining','fromdate','todate','participant','hparticipant','hour']
def replace(shape,m):
 if not getattr(shape,'has_text_frame',False): return
 for p in shape.text_frame.paragraphs:
  for r in p.runs:
   for a,b in m.items(): r.text=r.text.replace(a,b)
def make_pptx(d,out):
 prs=Presentation(TEMPLATE); m={f'<<{k}>>':str(d[k]) for k in FIELDS if k!='email'}
 for s in prs.slides:
  for sh in s.shapes: replace(sh,m)
 prs.save(out)
def pdf(p):
 o=p.parent/'pdf';o.mkdir(exist_ok=True)
 subprocess.run(['libreoffice','--headless','--convert-to','pdf','--outdir',str(o),str(p)],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
 return o/(p.stem+'.pdf')
def mail(to,p):
 msg=EmailMessage();msg['Subject']='CPD Certificate - PM SHRI Kendriya Vidyalaya Dongargarh';msg['From']=os.getenv('SMTP_FROM',os.environ['SMTP_USER']);msg['To']=to;msg.set_content('Dear Participant,\n\nPlease find attached your CPD certificate.\n\nRegards,\nPM SHRI Kendriya Vidyalaya Dongargarh');msg.add_attachment(p.read_bytes(),maintype='application',subtype='pdf',filename=p.name)
 with smtplib.SMTP(os.environ['SMTP_HOST'],int(os.getenv('SMTP_PORT','587'))) as s: s.starttls();s.login(os.environ['SMTP_USER'],os.environ['SMTP_PASSWORD']);s.send_message(msg)
@app.get('/api/health')
def health(): return {'ok':True}
@app.post('/api/generate')
def generate():
 d=request.get_json(silent=True) or {}; missing=[k for k in FIELDS if not str(d.get(k,'')).strip()]
 if missing:return jsonify(error='Missing fields: '+', '.join(missing)),400
 if not re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$',d['email']):return jsonify(error='Invalid email address'),400
 if not TEMPLATE.exists():return jsonify(error='PPTX template missing'),500
 try:
  p=OUT/(uuid.uuid4().hex+'.pptx');make_pptx(d,p);q=pdf(p);mail(d['email'],q)
  return jsonify(ok=True,download_url='/api/download/'+q.name,pptx_url='/api/download/'+p.name)
 except Exception as e:return jsonify(error='Generation failed: '+str(e)),500
@app.get('/api/download/<path:name>')
def download(name): return send_from_directory(OUT,name,as_attachment=True)
